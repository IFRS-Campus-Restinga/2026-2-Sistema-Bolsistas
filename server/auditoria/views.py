from io import BytesIO

from django.db.models import Q
from django.http import HttpResponse
from django.utils import timezone
from django.utils.dateparse import parse_date
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from rest_framework.exceptions import ValidationError
from rest_framework.pagination import PageNumberPagination
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from accounts.permissions import IsAdministrador

from .models import AcaoAuditoria, EntidadeAuditoria, LogAuditoria
from .serializers import LogAuditoriaSerializer


class PaginacaoAuditoria(PageNumberPagination):
    page_size = 20
    page_query_param = "pagina"
    page_size_query_param = "por_pagina"
    max_page_size = 100


def _opcoes(choices):
    return [{"valor": valor, "rotulo": rotulo} for valor, rotulo in choices.choices]


def _escolha_valida(request, parametro, choices):
    valor = request.query_params.get(parametro, "").strip()
    if valor and valor not in choices.values:
        raise ValidationError({parametro: f"Valor inválido: {valor}."})
    return valor


def _data(request, parametro, rotulo):
    texto = request.query_params.get(parametro)
    if not texto:
        return None
    data = parse_date(texto)
    if data is None:
        raise ValidationError({parametro: f"Informe a {rotulo} no formato AAAA-MM-DD."})
    return data


def _logs_filtrados(request):
    inicio = _data(request, "inicio", "data inicial")
    fim = _data(request, "fim", "data final")
    if inicio and fim and inicio > fim:
        raise ValidationError({"detail": "A data inicial não pode ser posterior à data final."})

    acao = _escolha_valida(request, "acao", AcaoAuditoria)
    entidade = _escolha_valida(request, "entidade", EntidadeAuditoria)
    usuario = request.query_params.get("usuario", "").strip()
    descricao = request.query_params.get("descricao", "").strip()

    logs = LogAuditoria.objects.all()
    if inicio:
        logs = logs.filter(data_hora__date__gte=inicio)
    if fim:
        logs = logs.filter(data_hora__date__lte=fim)
    if acao:
        logs = logs.filter(acao=acao)
    if entidade:
        logs = logs.filter(entidade=entidade)
    if usuario:
        logs = logs.filter(Q(usuario_nome__icontains=usuario) | Q(usuario_email__icontains=usuario))
    if descricao:
        logs = logs.filter(descricao__icontains=descricao)
    return logs


class LogAuditoriaListView(APIView):
    permission_classes = [IsAuthenticated, IsAdministrador]

    def get(self, request):
        paginador = PaginacaoAuditoria()
        pagina = paginador.paginate_queryset(_logs_filtrados(request), request, view=self)
        serializer = LogAuditoriaSerializer(pagina, many=True)
        return paginador.get_paginated_response(serializer.data)


COLUNAS_EXPORTACAO = [
    ("Data e hora", 20),
    ("Ação", 16),
    ("Usuário", 28),
    ("E-mail", 32),
    ("Entidade", 20),
    ("ID", 8),
    ("Descrição", 70),
]


def _planilha_de_logs(logs):
    planilha = Workbook()
    aba = planilha.active
    aba.title = "Auditoria"

    aba.append([titulo for titulo, _ in COLUNAS_EXPORTACAO])
    for celula in aba[1]:
        celula.font = Font(bold=True)
    for indice, (_, largura) in enumerate(COLUNAS_EXPORTACAO, start=1):
        aba.column_dimensions[get_column_letter(indice)].width = largura

    for log in logs:
        aba.append(
            [
                timezone.localtime(log.data_hora).replace(tzinfo=None),
                log.get_acao_display(),
                log.usuario_nome or "Sistema",
                log.usuario_email,
                log.get_entidade_display(),
                log.entidade_id,
                log.descricao,
            ]
        )
        aba.cell(row=aba.max_row, column=1).number_format = "dd/mm/yyyy hh:mm:ss"

    aba.freeze_panes = "A2"
    aba.auto_filter.ref = aba.dimensions
    return planilha


class ExportarLogsView(APIView):
    permission_classes = [IsAuthenticated, IsAdministrador]

    def get(self, request):
        conteudo = BytesIO()
        _planilha_de_logs(_logs_filtrados(request)).save(conteudo)

        resposta = HttpResponse(
            conteudo.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        hoje = timezone.localdate().isoformat()
        resposta["Content-Disposition"] = f'attachment; filename="auditoria-{hoje}.xlsx"'
        return resposta


class OpcoesAuditoriaView(APIView):
    permission_classes = [IsAuthenticated, IsAdministrador]

    def get(self, request):
        return Response(
            {
                "acoes": _opcoes(AcaoAuditoria),
                "entidades": _opcoes(EntidadeAuditoria),
                "hoje": timezone.localdate().isoformat(),
            }
        )
