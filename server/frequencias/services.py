from collections import Counter
from datetime import date

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from bolsas.models import StatusBolsa, TipoBolsa
from bolsistas.models import SituacaoRelatorio, StatusVinculo, VinculoBolsista
from bolsistas.regras import relatorio_do_vinculo

from .models import Frequencia, SolicitacaoFrequencia, StatusFrequencia
from .regras import (
    frequencia_pendente,
    meses_do_vinculo,
    motivo_bloqueio_lancamento,
    prazo_do_mes,
    status_do_mes,
)
from .serializers import FrequenciaSerializer, representar_mes


def vinculos_da_area(usuario):
    area = usuario.perfil_coordenador_area.tipo_area
    return VinculoBolsista.objects.filter(
        bolsa__tipo__in=[area, TipoBolsa.INDISSOCIAVEL]
    ).select_related(
        "aluno",
        "bolsa__edital",
        "bolsa__projeto__coordenador_projeto__usuario",
    )


def frequencias_indexadas(vinculos):
    """Frequências dos vínculos, por (vínculo, mês), em uma consulta só."""
    frequencias = Frequencia.objects.filter(vinculo__in=vinculos)
    return {(f.vinculo_id, f.mes_referencia): f for f in frequencias}


def solicitacoes_abertas_indexadas(vinculos):
    solicitacoes = SolicitacaoFrequencia.objects.filter(
        vinculo__in=vinculos, atendida_em__isnull=True
    ).select_related("solicitante")
    return {(s.vinculo_id, s.mes_referencia): s for s in solicitacoes}


def _formatar_data(data):
    return data.isoformat() if data else None


def situacao_da_area(usuario, ano, hoje=None):
    """Status de cada bolsista da área em cada mês do ano (US56)."""
    hoje = hoje or timezone.localdate()
    ultimo_mes = 12 if ano < hoje.year else hoje.month
    meses = (
        [date(ano, numero, 1) for numero in range(1, ultimo_mes + 1)] if ano <= hoje.year else []
    )

    vinculos = list(vinculos_da_area(usuario).order_by("bolsa__projeto__titulo", "aluno__nome"))
    frequencias = frequencias_indexadas(vinculos)
    totais = Counter()
    linhas = []

    for vinculo in vinculos:
        periodo = set(meses_do_vinculo(vinculo, hoje))
        if not periodo.intersection(meses):
            continue
        celulas = []
        for mes in meses:
            if mes not in periodo:
                celulas.append(None)
                continue
            frequencia = frequencias.get((vinculo.id, mes))
            status = status_do_mes(frequencia)
            totais[status] += 1
            celulas.append(
                {
                    "status": status,
                    "pendente": frequencia_pendente(vinculo, mes, frequencia, hoje),
                    "fora_do_prazo": bool(frequencia and frequencia.fora_do_prazo),
                }
            )
        linhas.append(
            {
                "vinculo_id": vinculo.id,
                "aluno_nome": vinculo.aluno.nome,
                "status_vinculo": vinculo.status,
                "projeto_titulo": vinculo.bolsa.projeto.titulo,
                "coordenador_projeto_nome": vinculo.bolsa.projeto.coordenador_projeto.usuario.nome,
                "edital_ano_codigo": vinculo.bolsa.edital.ano_codigo,
                "meses": celulas,
            }
        )

    return {
        "ano": ano,
        "meses": [representar_mes(mes) for mes in meses],
        "totais": {status: totais[status] for status in StatusFrequencia.values},
        "bolsistas": linhas,
    }


def _relatorio_pendente(vinculo, hoje):
    """Relatório final não enviado e com o prazo do edital vencido; None quando em dia."""
    prazo = vinculo.bolsa.edital.data_entrega_relatorios
    if vinculo.status != StatusVinculo.ATIVO or prazo is None or hoje <= prazo:
        return None
    relatorio = relatorio_do_vinculo(vinculo)
    if relatorio is None:
        situacao, rotulo = "NAO_SOLICITADO", "Não solicitado"
    elif relatorio.situacao == SituacaoRelatorio.SOLICITADO:
        situacao, rotulo = relatorio.situacao, "Não enviado"
    else:
        return None
    return {
        "vinculo_id": vinculo.id,
        "aluno_nome": vinculo.aluno.nome,
        "situacao": situacao,
        "situacao_display": rotulo,
        "prazo": _formatar_data(prazo),
    }


def pendencias_da_area(usuario, hoje=None):
    """Frequências e relatórios em atraso, agrupados por bolsa e coordenador (US58)."""
    hoje = hoje or timezone.localdate()
    vinculos = list(
        vinculos_da_area(usuario)
        .exclude(bolsa__status=StatusBolsa.ENCERRADA)
        .select_related("relatorio_final")
        .order_by("bolsa__projeto__coordenador_projeto__usuario__nome", "bolsa__projeto__titulo")
    )
    frequencias = frequencias_indexadas(vinculos)
    solicitacoes = solicitacoes_abertas_indexadas(vinculos)
    grupos = {}

    for vinculo in vinculos:
        bolsa = vinculo.bolsa
        pendentes = [
            mes
            for mes in meses_do_vinculo(vinculo, hoje)
            if frequencia_pendente(vinculo, mes, frequencias.get((vinculo.id, mes)), hoje)
        ]
        relatorio = _relatorio_pendente(vinculo, hoje)
        if not pendentes and relatorio is None:
            continue

        grupo = grupos.setdefault(
            bolsa.id,
            {
                "bolsa_id": bolsa.id,
                "projeto_titulo": bolsa.projeto.titulo,
                "coordenador_projeto_nome": bolsa.projeto.coordenador_projeto.usuario.nome,
                "edital_ano_codigo": bolsa.edital.ano_codigo,
                "tipo_display": bolsa.get_tipo_display(),
                "frequencias": [],
                "relatorios": [],
            },
        )
        for mes in pendentes:
            solicitacao = solicitacoes.get((vinculo.id, mes))
            grupo["frequencias"].append(
                {
                    "vinculo_id": vinculo.id,
                    "aluno_nome": vinculo.aluno.nome,
                    "mes_referencia": representar_mes(mes),
                    "prazo": _formatar_data(prazo_do_mes(bolsa.edital, mes)),
                    "solicitada_em": solicitacao.criada_em if solicitacao else None,
                    "lida_em": solicitacao.lida_em if solicitacao else None,
                }
            )
        if relatorio:
            grupo["relatorios"].append(relatorio)

    return list(grupos.values())


def mes_do_coordenador(vinculo, mes, frequencia, solicitacao, hoje=None):
    """Linha do formulário de frequência do coordenador do projeto para um mês."""
    motivo = motivo_bloqueio_lancamento(vinculo, mes, frequencia, hoje)
    return {
        "vinculo_id": vinculo.id,
        "aluno_nome": vinculo.aluno.nome,
        "aluno_email": vinculo.aluno.email,
        "status_vinculo": vinculo.status,
        "projeto_titulo": vinculo.bolsa.projeto.titulo,
        "edital_ano_codigo": vinculo.bolsa.edital.ano_codigo,
        "bolsa_tipo_display": vinculo.bolsa.get_tipo_display(),
        "mes_referencia": representar_mes(mes),
        "prazo": _formatar_data(prazo_do_mes(vinculo.bolsa.edital, mes)),
        "status": status_do_mes(frequencia),
        "pendente": frequencia_pendente(vinculo, mes, frequencia, hoje),
        "frequencia": FrequenciaSerializer(frequencia).data if frequencia else None,
        "solicitada_em": solicitacao.criada_em if solicitacao else None,
        "pode_lancar": not motivo,
        "motivo_bloqueio": motivo,
    }


def historico_do_aluno(usuario, hoje=None):
    """Mês a mês de cada bolsa do aluno, com as observações e os pareceres (US60)."""
    hoje = hoje or timezone.localdate()
    vinculos = list(
        VinculoBolsista.objects.filter(aluno=usuario)
        .select_related(
            "bolsa__edital",
            "bolsa__projeto__coordenador_projeto__usuario",
            "relatorio_final",
        )
        .order_by("-criado_em", "-id")
    )
    frequencias = frequencias_indexadas(vinculos)
    historico = []

    for vinculo in vinculos:
        meses = []
        for mes in reversed(meses_do_vinculo(vinculo, hoje)):
            frequencia = frequencias.get((vinculo.id, mes))
            status = status_do_mes(frequencia)
            meses.append(
                {
                    "mes_referencia": representar_mes(mes),
                    "status": status,
                    "status_display": status.label,
                    "observacoes": frequencia.observacoes if frequencia else "",
                    "lancada_em": frequencia.lancada_em if frequencia else None,
                    "fora_do_prazo": bool(frequencia and frequencia.fora_do_prazo),
                    "status_pagamento": frequencia.status_pagamento if frequencia else None,
                    "status_pagamento_display": (
                        frequencia.get_status_pagamento_display() if frequencia else ""
                    ),
                    "justificativa_pagamento": (
                        frequencia.justificativa_pagamento if frequencia else ""
                    ),
                }
            )
        relatorio = relatorio_do_vinculo(vinculo)
        historico.append(
            {
                "vinculo_id": vinculo.id,
                "status_vinculo": vinculo.status,
                "projeto_titulo": vinculo.bolsa.projeto.titulo,
                "coordenador_nome": vinculo.bolsa.projeto.coordenador_projeto.usuario.nome,
                "edital_ano_codigo": vinculo.bolsa.edital.ano_codigo,
                "edital_nome": vinculo.bolsa.edital.nome,
                "relatorio": (
                    {
                        "situacao": relatorio.situacao,
                        "situacao_display": relatorio.get_situacao_display(),
                        "justificativa": relatorio.justificativa,
                    }
                    if relatorio
                    else None
                ),
                "meses": meses,
            }
        )
    return historico


COLUNAS_EXPORTACAO = [
    ("Projeto", 36),
    ("Edital", 12),
    ("Bolsista", 30),
    ("E-mail", 30),
    ("Vínculo", 12),
    ("Mês", 10),
    ("Frequência", 16),
    ("Prazo de lançamento", 18),
    ("Lançada em", 18),
    ("Fora do prazo", 13),
    ("Observações", 40),
    ("Pagamento", 13),
    ("Justificativa do pagamento", 40),
    ("Relatório final", 22),
]


def planilha_de_acompanhamento(vinculos, hoje=None):
    """Planilha com a frequência e os status de cada bolsista, mês a mês (US59)."""
    hoje = hoje or timezone.localdate()
    vinculos = list(vinculos)
    frequencias = frequencias_indexadas(vinculos)

    planilha = Workbook()
    aba = planilha.active
    aba.title = "Acompanhamento"
    aba.append([titulo for titulo, _ in COLUNAS_EXPORTACAO])
    for celula in aba[1]:
        celula.font = Font(bold=True)
    for indice, (_, largura) in enumerate(COLUNAS_EXPORTACAO, start=1):
        aba.column_dimensions[get_column_letter(indice)].width = largura

    for vinculo in vinculos:
        relatorio = relatorio_do_vinculo(vinculo)
        situacao_relatorio = relatorio.get_situacao_display() if relatorio else "Não solicitado"
        for mes in meses_do_vinculo(vinculo, hoje):
            frequencia = frequencias.get((vinculo.id, mes))
            aba.append(
                [
                    vinculo.bolsa.projeto.titulo,
                    vinculo.bolsa.edital.ano_codigo,
                    vinculo.aluno.nome,
                    vinculo.aluno.email or "",
                    vinculo.get_status_display(),
                    f"{mes:%m/%Y}",
                    status_do_mes(frequencia).label,
                    prazo_do_mes(vinculo.bolsa.edital, mes),
                    (
                        timezone.localtime(frequencia.lancada_em).replace(tzinfo=None)
                        if frequencia
                        else None
                    ),
                    ("Sim" if frequencia.fora_do_prazo else "Não") if frequencia else "",
                    frequencia.observacoes if frequencia else "",
                    frequencia.get_status_pagamento_display() if frequencia else "",
                    frequencia.justificativa_pagamento if frequencia else "",
                    situacao_relatorio,
                ]
            )
            aba.cell(row=aba.max_row, column=8).number_format = "dd/mm/yyyy"
            aba.cell(row=aba.max_row, column=9).number_format = "dd/mm/yyyy hh:mm"

    aba.freeze_panes = "A2"
    aba.auto_filter.ref = aba.dimensions
    return planilha
