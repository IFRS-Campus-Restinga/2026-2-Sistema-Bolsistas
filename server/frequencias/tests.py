from datetime import date, datetime, timedelta
from io import BytesIO

from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APITestCase

from accounts.models import CoordenadorArea, CoordenadorProjeto, TipoArea, Usuario
from auditoria.models import EntidadeAuditoria, LogAuditoria
from bolsas.models import Bolsa, StatusBolsa
from bolsistas.models import OrigemVinculo, StatusVinculo, VinculoBolsista
from editais.models import Edital
from inscricoes.models import Inscricao, StatusInscricao
from projetos.models import Projeto
from selecao.models import IndicacaoBolsista

from .models import Frequencia, SolicitacaoFrequencia, StatusFrequencia, StatusPagamento
from .regras import meses_do_vinculo, prazo_do_mes, primeiro_dia, proximo_mes


def meses_atras(quantidade):
    mes = primeiro_dia(timezone.localdate())
    for _ in range(quantidade):
        mes = primeiro_dia(mes - timedelta(days=1))
    return mes


class CenarioFrequencia(APITestCase):
    def setUp(self):
        self.coordenador = Usuario.objects.create_user(
            username="coordenador",
            email="coordenador@example.com",
            nome="Coordenador",
            role=Usuario.Role.COORDENADOR_PROJETO,
        )
        self.projeto = Projeto.objects.create(
            titulo="Projeto",
            coordenador_projeto=CoordenadorProjeto.objects.create(usuario=self.coordenador),
        )
        self.edital = Edital.objects.create(
            nome="Edital",
            ano_codigo="2026-001",
            link_documento_oficial="https://example.com/edital",
            status=Edital.Status.EM_VIGOR,
            dia_limite_frequencia=1,
        )
        self.bolsa = Bolsa.objects.create(
            projeto=self.projeto,
            edital=self.edital,
            tipo="ENSINO",
            modalidade="BICT",
            carga_horaria_semanal=12,
            valor_mensal=700,
            quantidade_vagas=3,
            status=StatusBolsa.PREENCHIDA,
        )
        self.area = Usuario.objects.create_user(
            username="area",
            email="area@example.com",
            nome="Área",
            role=Usuario.Role.COORDENADOR_AREA,
        )
        CoordenadorArea.objects.create(usuario=self.area, tipo_area=TipoArea.ENSINO)
        self.vinculo = self.criar_vinculo("ana")

    def criar_vinculo(self, nome, meses_de_bolsa=3):
        aluno = Usuario.objects.create_user(
            username=nome, email=f"{nome}@example.com", nome=nome, role=Usuario.Role.ALUNO
        )
        inscricao = Inscricao.objects.create(
            aluno=aluno, bolsa=self.bolsa, status=StatusInscricao.HOMOLOGADA
        )
        vinculo = VinculoBolsista.objects.create(
            bolsa=self.bolsa,
            aluno=aluno,
            origem=OrigemVinculo.SELECAO,
            indicacao=IndicacaoBolsista.objects.create(
                bolsa=self.bolsa, inscricao=inscricao, responsavel=self.coordenador
            ),
        )
        inicio = timezone.make_aware(
            datetime.combine(meses_atras(meses_de_bolsa), datetime.min.time())
        )
        VinculoBolsista.objects.filter(pk=vinculo.pk).update(criado_em=inicio)
        vinculo.refresh_from_db()
        return vinculo

    def lancar(self, mes, registrada=True, observacoes="", vinculo=None):
        self.client.force_authenticate(self.coordenador)
        return self.client.post(
            reverse("frequencias:lancar"),
            {
                "vinculo": (vinculo or self.vinculo).id,
                "mes_referencia": f"{mes:%Y-%m}",
                "registrada": registrada,
                "observacoes": observacoes,
            },
            format="json",
        )


class RegrasDePrazoTests(CenarioFrequencia):
    def test_prazo_e_o_dia_limite_do_mes_seguinte(self):
        self.edital.dia_limite_frequencia = 10
        self.assertEqual(prazo_do_mes(self.edital, date(2026, 10, 1)), date(2026, 11, 10))
        self.assertEqual(prazo_do_mes(self.edital, date(2026, 12, 1)), date(2027, 1, 10))

    def test_dia_limite_acima_do_fim_do_mes_vale_o_ultimo_dia(self):
        self.edital.dia_limite_frequencia = 31
        self.assertEqual(prazo_do_mes(self.edital, date(2026, 1, 1)), date(2026, 2, 28))

    def test_edital_sem_dia_limite_nao_tem_prazo(self):
        self.edital.dia_limite_frequencia = None
        self.assertIsNone(prazo_do_mes(self.edital, date(2026, 10, 1)))

    def test_periodo_vai_do_ingresso_ate_o_desligamento(self):
        self.vinculo.status = StatusVinculo.DESLIGADO
        self.vinculo.desligado_em = timezone.now() - timedelta(days=40)
        meses = meses_do_vinculo(self.vinculo)
        self.assertEqual(meses[0], meses_atras(3))
        self.assertEqual(meses[-1], primeiro_dia(timezone.localdate(self.vinculo.desligado_em)))

    def test_periodo_respeita_o_fim_da_vigencia(self):
        self.vinculo.bolsa.data_fim_vigencia = meses_atras(2) + timedelta(days=5)
        self.assertEqual(meses_do_vinculo(self.vinculo)[-1], meses_atras(2))


class LancamentoFrequenciaTests(CenarioFrequencia):
    def test_coordenador_lanca_frequencia_no_prazo(self):
        resposta = self.lancar(meses_atras(0))

        self.assertEqual(resposta.status_code, 201, resposta.data)
        frequencia = Frequencia.objects.get()
        self.assertTrue(frequencia.registrada)
        self.assertFalse(frequencia.fora_do_prazo)
        self.assertEqual(frequencia.status_pagamento, StatusPagamento.PENDENTE)
        self.assertTrue(LogAuditoria.objects.filter(entidade=EntidadeAuditoria.FREQUENCIA).exists())

    def test_lancamento_depois_do_prazo_e_sinalizado(self):
        resposta = self.lancar(meses_atras(2))

        self.assertEqual(resposta.status_code, 201, resposta.data)
        self.assertTrue(resposta.data["fora_do_prazo"])

    def test_relancar_corrige_sem_perder_o_atraso(self):
        self.lancar(meses_atras(2))
        resposta = self.lancar(meses_atras(2), registrada=False, observacoes="Faltou.")

        self.assertEqual(resposta.status_code, 200, resposta.data)
        frequencia = Frequencia.objects.get()
        self.assertFalse(frequencia.registrada)
        self.assertTrue(frequencia.fora_do_prazo)

    def test_frequencia_nao_registrada_exige_observacao(self):
        resposta = self.lancar(meses_atras(0), registrada=False)

        self.assertEqual(resposta.status_code, 400)
        self.assertIn("observacoes", resposta.data)

    def test_nao_lanca_mes_fora_do_periodo(self):
        resposta = self.lancar(proximo_mes(meses_atras(0)))

        self.assertEqual(resposta.status_code, 400)
        self.assertFalse(Frequencia.objects.exists())

    def test_nao_lanca_para_bolsista_de_outro_coordenador(self):
        outro = Usuario.objects.create_user(
            username="outro", email="outro@example.com", role=Usuario.Role.COORDENADOR_PROJETO
        )
        self.client.force_authenticate(outro)
        resposta = self.client.post(
            reverse("frequencias:lancar"),
            {
                "vinculo": self.vinculo.id,
                "mes_referencia": f"{meses_atras(0):%Y-%m}",
                "registrada": True,
            },
            format="json",
        )
        self.assertEqual(resposta.status_code, 404)

    def test_bolsa_encerrada_nao_aceita_lancamento(self):
        Bolsa.objects.filter(pk=self.bolsa.pk).update(status=StatusBolsa.ENCERRADA)
        self.assertEqual(self.lancar(meses_atras(0)).status_code, 400)

    def test_lista_do_mes_mostra_status_e_pendencia(self):
        self.client.force_authenticate(self.coordenador)
        resposta = self.client.get(
            reverse("frequencias:do-mes"), {"mes": f"{meses_atras(2):%Y-%m}"}
        )

        self.assertEqual(resposta.status_code, 200)
        linha = resposta.data["bolsistas"][0]
        self.assertEqual(linha["status"], StatusFrequencia.NAO_INFORMADO)
        self.assertTrue(linha["pendente"])
        self.assertTrue(linha["pode_lancar"])


class PagamentoTests(CenarioFrequencia):
    def decidir(self, frequencia_id, **dados):
        self.client.force_authenticate(self.coordenador)
        return self.client.post(
            reverse("frequencias:pagamento", kwargs={"pk": frequencia_id}), dados, format="json"
        )

    def test_autoriza_pagamento_de_frequencia_registrada(self):
        frequencia_id = self.lancar(meses_atras(0)).data["id"]
        resposta = self.decidir(frequencia_id, autorizar=True)

        self.assertEqual(resposta.status_code, 200, resposta.data)
        self.assertEqual(resposta.data["status_pagamento"], StatusPagamento.AUTORIZADO)

    def test_nao_autoriza_mes_sem_frequencia_registrada(self):
        frequencia_id = self.lancar(meses_atras(0), registrada=False, observacoes="Faltou.").data[
            "id"
        ]
        self.assertEqual(self.decidir(frequencia_id, autorizar=True).status_code, 400)

    def test_negar_exige_justificativa(self):
        frequencia_id = self.lancar(meses_atras(0)).data["id"]

        self.assertEqual(self.decidir(frequencia_id, autorizar=False).status_code, 400)
        resposta = self.decidir(frequencia_id, autorizar=False, justificativa="Sem relatório.")
        self.assertEqual(resposta.data["status_pagamento"], StatusPagamento.NEGADO)

    def test_decisao_trava_a_frequencia(self):
        frequencia_id = self.lancar(meses_atras(0)).data["id"]
        self.decidir(frequencia_id, autorizar=True)

        self.assertEqual(self.decidir(frequencia_id, autorizar=True).status_code, 400)
        self.assertEqual(self.lancar(meses_atras(0)).status_code, 400)


class SolicitacaoFrequenciaTests(CenarioFrequencia):
    def solicitar(self, mes, usuario=None):
        self.client.force_authenticate(usuario or self.area)
        return self.client.post(
            reverse("frequencias:solicitar"),
            {"vinculo": self.vinculo.id, "mes_referencia": f"{mes:%Y-%m}", "mensagem": "Lançar."},
            format="json",
        )

    def test_area_solicita_frequencia_em_atraso_e_coordenador_e_notificado(self):
        resposta = self.solicitar(meses_atras(2))
        self.assertEqual(resposta.status_code, 201, resposta.data)

        self.client.force_authenticate(self.coordenador)
        recebidas = self.client.get(reverse("frequencias:solicitacoes-recebidas"))
        self.assertEqual(len(recebidas.data), 1)
        self.assertIsNone(recebidas.data[0]["lida_em"])

        ciencia = self.client.post(
            reverse("frequencias:ciencia-solicitacao", kwargs={"pk": resposta.data["id"]})
        )
        self.assertIsNotNone(ciencia.data["lida_em"])

    def test_lancamento_atende_a_solicitacao(self):
        self.solicitar(meses_atras(2))
        self.lancar(meses_atras(2))

        self.assertIsNotNone(SolicitacaoFrequencia.objects.get().atendida_em)

    def test_nao_solicita_mes_ainda_no_prazo(self):
        self.assertEqual(self.solicitar(meses_atras(0)).status_code, 400)

    def test_nao_duplica_solicitacao_aberta(self):
        self.solicitar(meses_atras(2))
        self.assertEqual(self.solicitar(meses_atras(2)).status_code, 400)

    def test_area_de_outro_tipo_nao_solicita(self):
        pesquisa = Usuario.objects.create_user(
            username="pesquisa", email="pesquisa@example.com", role=Usuario.Role.COORDENADOR_AREA
        )
        CoordenadorArea.objects.create(usuario=pesquisa, tipo_area=TipoArea.PESQUISA)
        self.assertEqual(self.solicitar(meses_atras(2), pesquisa).status_code, 404)


class VisaoDaAreaTests(CenarioFrequencia):
    def status_no_painel(self, mes):
        self.client.force_authenticate(self.area)
        resposta = self.client.get(reverse("frequencias:situacao-area"), {"ano": mes.year})
        self.assertEqual(resposta.status_code, 200)
        celulas = dict(
            zip(resposta.data["meses"], resposta.data["bolsistas"][0]["meses"], strict=True)
        )
        return celulas[f"{mes:%Y-%m}"]

    def test_painel_mostra_os_tres_status(self):
        self.lancar(meses_atras(3))
        self.lancar(meses_atras(2), registrada=False, observacoes="Faltou.")

        self.assertEqual(
            self.status_no_painel(meses_atras(3))["status"], StatusFrequencia.REGISTRADA
        )
        self.assertEqual(
            self.status_no_painel(meses_atras(2))["status"], StatusFrequencia.NAO_REGISTRADA
        )
        celula = self.status_no_painel(meses_atras(1))
        self.assertEqual(celula["status"], StatusFrequencia.NAO_INFORMADO)

    def test_pendencias_agrupam_frequencias_e_relatorios_por_bolsa(self):
        self.edital.data_entrega_relatorios = timezone.localdate() - timedelta(days=1)
        self.edital.save()

        self.client.force_authenticate(self.area)
        resposta = self.client.get(reverse("frequencias:pendencias-area"))

        self.assertEqual(resposta.status_code, 200)
        self.assertEqual(len(resposta.data), 1)
        grupo = resposta.data[0]
        self.assertEqual(grupo["coordenador_projeto_nome"], "Coordenador")
        meses_pendentes = {item["mes_referencia"] for item in grupo["frequencias"]}
        self.assertIn(f"{meses_atras(2):%Y-%m}", meses_pendentes)
        self.assertNotIn(f"{meses_atras(0):%Y-%m}", meses_pendentes)
        self.assertEqual(grupo["relatorios"][0]["situacao"], "NAO_SOLICITADO")


class ExportacaoEHistoricoTests(CenarioFrequencia):
    def test_exporta_planilha_com_uma_linha_por_mes(self):
        self.lancar(meses_atras(1))

        resposta = self.client.get(reverse("frequencias:exportar"))

        self.assertEqual(resposta.status_code, 200)
        self.assertIn(".xlsx", resposta["Content-Disposition"])
        aba = load_workbook(BytesIO(resposta.content)).active
        self.assertEqual(aba.max_row, 1 + len(meses_do_vinculo(self.vinculo)))

    def test_aluno_ve_o_proprio_historico(self):
        self.lancar(meses_atras(1), registrada=False, observacoes="Faltou duas semanas.")
        outro = self.criar_vinculo("bia")

        self.client.force_authenticate(self.vinculo.aluno)
        resposta = self.client.get(reverse("frequencias:minhas"))

        self.assertEqual(resposta.status_code, 200)
        self.assertEqual([item["vinculo_id"] for item in resposta.data], [self.vinculo.id])
        self.assertNotEqual(outro.id, self.vinculo.id)
        mes = next(
            item
            for item in resposta.data[0]["meses"]
            if item["mes_referencia"] == f"{meses_atras(1):%Y-%m}"
        )
        self.assertEqual(mes["status"], StatusFrequencia.NAO_REGISTRADA)
        self.assertEqual(mes["observacoes"], "Faltou duas semanas.")

    def test_somente_aluno_acessa_historico(self):
        self.client.force_authenticate(self.coordenador)
        self.assertEqual(self.client.get(reverse("frequencias:minhas")).status_code, 403)
